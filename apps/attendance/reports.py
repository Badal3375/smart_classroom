"""Attendance register views (daily/weekly/monthly/semester) and CSV / Excel / PDF exports."""
import csv
import io
from datetime import timedelta

from django.http import HttpResponse
from django.shortcuts import render
from django.core.paginator import Paginator
from django.utils import timezone

from apps.accounts.decorators import teacher_required
from apps.academics.models import Subject
from apps.core.audit import log_action
from .models import Attendance

PERIODS = {"daily": 1, "weekly": 7, "monthly": 30, "semester": 183}


def filtered_records(request):
    qs = Attendance.objects.select_related("student__user", "subject", "classroom")
    u = request.user
    if not (u.is_superuser or u.role == "ADMIN"):
        qs = qs.filter(session__teacher__user=u)
    period = request.GET.get("period", "weekly")
    today = timezone.localdate()
    d_from, d_to = request.GET.get("from"), request.GET.get("to")
    if d_from and d_to:
        qs = qs.filter(date__range=(d_from, d_to))
    else:
        qs = qs.filter(date__gte=today - timedelta(days=PERIODS.get(period, 7) - 1))
    if request.GET.get("subject", "").isdigit():
        qs = qs.filter(subject_id=request.GET["subject"])
    if request.GET.get("status") in ("P", "A", "L", "E"):
        qs = qs.filter(status=request.GET["status"])
    q = request.GET.get("q", "").strip()
    if q:
        from django.db.models import Q
        qs = qs.filter(Q(student__roll_no__icontains=q) | Q(student__user__first_name__icontains=q) | Q(student__user__last_name__icontains=q))
    return qs.order_by("-date", "student__roll_no"), period


@teacher_required
def register(request):
    qs, period = filtered_records(request)
    page = Paginator(qs, 20).get_page(request.GET.get("page"))
    subjects = Subject.objects.all()
    if not (request.user.is_superuser or request.user.role == "ADMIN"):
        subjects = subjects.filter(teachers__user=request.user)
    return render(request, "teacher/register.html", {"page": page, "period": period, "subjects": subjects,
                                                     "periods": PERIODS, "statuses": Attendance.STATUSES, "get": request.GET})


HEAD = ["Date", "Time", "Roll No", "Student", "Subject", "Classroom", "Status", "Confidence %", "Method", "Session"]


def _rows(qs):
    for a in qs:
        yield [a.date.isoformat(), a.time.strftime("%H:%M:%S") if a.time else "", a.student.roll_no, a.student.name,
               a.subject.code, a.classroom.name, a.get_status_display(), round(a.confidence * 100, 1), a.method, a.session_id]


@teacher_required
def export(request, fmt):
    qs, period = filtered_records(request)
    log_action(request.user, "REPORT_EXPORT", fmt, request.GET.urlencode(), request)
    name = f"attendance_{timezone.localdate():%Y%m%d}"
    if fmt == "csv":
        resp = HttpResponse(content_type="text/csv")
        resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
        w = csv.writer(resp)
        w.writerow(HEAD)
        w.writerows(_rows(qs))
        return resp
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"
        ws.append(HEAD)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4F46E5")
        for r in _rows(qs):
            ws.append(r)
        for col, wd in zip("ABCDEFGHIJ", [12, 10, 12, 24, 10, 12, 10, 13, 10, 9]):
            ws.column_dimensions[col].width = wd
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
        return resp
    if fmt == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
        st = getSampleStyleSheet()
        data = [HEAD[:8]] + [r[:8] for r in _rows(qs[:1500])]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                               ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")])]))
        doc.build([Paragraph("Attendance Report", st["Title"]),
                   Paragraph(f"Generated {timezone.localtime():%d %b %Y %H:%M} | Period: {period} | Records: {min(qs.count(), 1500)}", st["Normal"]),
                   Spacer(1, 10), t])
        resp = HttpResponse(buf.getvalue(), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="{name}.pdf"'
        return resp
    return HttpResponse(status=404)
